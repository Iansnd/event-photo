#!/usr/bin/env python3
"""Merge regional research JSON into leads CSV/XLSX. Usage: merge_leads.py <regions_dir> <out_dir>"""
import csv, glob, json, os, sys, re
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

src, out = sys.argv[1], sys.argv[2]
os.makedirs(out, exist_ok=True)

LEAD_COLS = ["org_name","org_type","country","countries","hq_city","hq_address","website","num_settings","num_children","revenue",
             "ownership","decision_maker_name","decision_maker_title","decision_maker_linkedin","other_contacts",
             "general_email","phone","background","confirmation_status","notes","lead_score","sources","region_file"]
SRC_COLS = ["country","source_name","url","what_it_contains","how_to_extract","cost","region_file"]

def flat(v):
    if v is None: return ""
    if isinstance(v, list):
        return "; ".join(flat(x) for x in v)
    if isinstance(v, dict):
        return "; ".join(f"{k}: {flat(x)}" for k, x in v.items() if x not in (None, "", []))
    return str(v).strip()

leads, sources = [], []
for fp in sorted(glob.glob(os.path.join(src, "*.json"))):
    base = os.path.basename(fp)
    try:
        data = json.load(open(fp, encoding="utf-8"))
    except Exception as e:
        print("SKIP", base, e); continue
    if isinstance(data, dict):
        data = data.get("leads") or data.get("sources") or data.get("data") or list(data.values())[0]
    for d in data:
        if not isinstance(d, dict): continue
        d = {k: flat(v) for k, v in d.items()}
        d["region_file"] = base
        if base.endswith("_sources.json") or ("source_name" in d and "org_name" not in d):
            sources.append({c: d.get(c, "") for c in SRC_COLS})
        else:
            if "country" not in d or not d["country"]:
                d["country"] = d.get("hq_country", "")
            if not d.get("confirmation_status"):
                d["confirmation_status"] = d.get("dm_confirmation") or d.get("verification") or ""
            leads.append({c: d.get(c, "") for c in LEAD_COLS})

# dedupe by normalised org name + country
def key(d): return (re.sub(r"[^a-z0-9]", "", d["org_name"].lower()), d["country"][:2].upper())
seen, merged = {}, []
for d in leads:
    k = key(d)
    if k in seen:
        m = seen[k]
        for c in LEAD_COLS:
            if not m[c] and d[c]: m[c] = d[c]
            elif c in ("background","sources","other_contacts") and d[c] and d[c] not in m[c]:
                m[c] = (m[c] + " | " + d[c]).strip(" |")
    else:
        seen[k] = d; merged.append(d)

def score(d):
    try: return -float(d["lead_score"] or 0)
    except: return 0
merged.sort(key=lambda d: (score(d), d["country"], d["org_name"]))

def write_csv(path, rows, cols):
    with open(path, "w", newline="", encoding="utf-8") as f:
        w = csv.DictWriter(f, fieldnames=cols); w.writeheader(); w.writerows(rows)

write_csv(os.path.join(out, "kinderlee_leads_europe.csv"), merged, LEAD_COLS)
write_csv(os.path.join(out, "free_lead_sources.csv"), sources, SRC_COLS)

# decision-maker sheet
dm = [d for d in merged if d["decision_maker_name"] and d["decision_maker_name"].lower() not in ("null","none","not confirmed")]
DM_COLS = ["decision_maker_name","decision_maker_title","org_name","country","decision_maker_linkedin","general_email","phone","website","lead_score"]
write_csv(os.path.join(out, "decision_makers.csv"), [{c: d[c] for c in DM_COLS} for d in dm], DM_COLS)

wb = Workbook()
def sheet(ws, rows, cols, widths):
    ws.append(cols)
    for c in ws[1]:
        c.font = Font(bold=True, color="FFFFFF"); c.fill = PatternFill("solid", fgColor="1F4E79"); c.alignment = Alignment(wrap_text=True, vertical="top")
    for r in rows: ws.append([r.get(c, "") for c in cols])
    for i, c in enumerate(cols, 1):
        ws.column_dimensions[get_column_letter(i)].width = widths.get(c, 18)
    for row in ws.iter_rows(min_row=2):
        for c in row: c.alignment = Alignment(wrap_text=True, vertical="top")
    ws.freeze_panes = "A2"; ws.auto_filter.ref = ws.dimensions
ws = wb.active; ws.title = "Leads"
sheet(ws, merged, LEAD_COLS, {"org_name":30,"background":80,"notes":40,"other_contacts":45,"sources":50,"hq_address":30,"ownership":28,"decision_maker_name":24,"decision_maker_linkedin":40,"website":30})
sheet(wb.create_sheet("Decision makers"), [{c: d[c] for c in DM_COLS} for d in dm], DM_COLS, {"decision_maker_name":26,"org_name":30,"decision_maker_linkedin":45,"website":30})
sheet(wb.create_sheet("Free data sources"), sources, SRC_COLS, {"source_name":32,"url":45,"what_it_contains":60,"how_to_extract":60})
# per-country summary
from collections import Counter
cnt = Counter(d["country"] for d in merged); dmc = Counter(d["country"] for d in dm)
sheet(wb.create_sheet("Summary"), [{"country":k,"leads":v,"with_decision_maker":dmc.get(k,0)} for k,v in sorted(cnt.items())], ["country","leads","with_decision_maker"], {})
wb.save(os.path.join(out, "kinderlee_leads_europe.xlsx"))
print(f"leads={len(merged)} decision_makers={len(dm)} sources={len(sources)} countries={len(cnt)}")
for k, v in sorted(cnt.items()): print(f"  {k}: {v} ({dmc.get(k,0)} with DM)")
